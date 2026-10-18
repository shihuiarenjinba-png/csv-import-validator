"""Read column names safely for the beginner Key picker.

This is read-only, never edits customer CSV/XLSX and never silently deduplicates
headers. The calling GUI loads it on a worker thread.
"""
from __future__ import annotations

import csv
from pathlib import Path

from audit_review_gate import input_revision


def read_header(path: str) -> list[str]:
    p = Path(path).expanduser()
    if not p.is_file() or p.is_symlink():
        raise ValueError("存在する通常のCSV/Excelファイルを指定してください。")
    if p.stat().st_size == 0 or p.stat().st_size > 100 * 1024 * 1024:
        raise ValueError("ファイルサイズは1バイト以上100MiB以下にしてください。")
    if p.suffix.lower() == ".csv":
        names = None
        for enc in ("utf-8-sig", "utf-8", "cp932", "shift_jis"):
            try:
                with p.open("r", encoding=enc, newline="") as stream:
                    names = next(csv.reader(stream, strict=True), [])
                break
            except UnicodeDecodeError:
                continue
        if names is None:
            raise ValueError("CSV文字コードを読み取れません。")
    elif p.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        workbook = load_workbook(p, read_only=True, data_only=True, keep_links=False)
        try:
            sheet = workbook.worksheets[0]
            cells = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            names = ["" if v is None else str(v) for v in cells]
        finally:
            workbook.close()
    else:
        raise ValueError("列の選択にはCSV/XLSX/XLSMを使ってください。")
    if not names or any(not n.strip() for n in names):
        raise ValueError("列名のない列があります。ファイルの1行目を確認してください。")
    if len(set(names)) != len(names):
        raise ValueError("同じ列名が複数あります。列名を修正してください。")
    return names


def common_key_candidates(baseline: str, current: str) -> list[str]:
    baseline_revision = input_revision(baseline, current, [])
    b = read_header(baseline)
    c = read_header(current)
    if input_revision(baseline, current, []) != baseline_revision:
        raise ValueError("列名を読み込んでいる間に入力が変わりました。再実行してください。")
    common = [col for col in b if col in c and "," not in col]
    if not common:
        raise ValueError("両方のファイルに共通する列がありません。比較設定を確認してください。")
    return common
